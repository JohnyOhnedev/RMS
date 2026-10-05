"""Resume text extraction for PDF, DOCX, TXT and image files.

Uses PyMuPDF as the primary PDF engine (fast, layout-aware) with pypdf as a
fallback, python-docx for Word documents, and optional Tesseract OCR for
scanned/image resumes.
"""

from __future__ import annotations

import logging
import os
import tempfile

from django.conf import settings

from .exceptions import EmptyResumeError, UnsupportedFileType

logger = logging.getLogger(__name__)


def _extract_pdf(path: str) -> str:
    """Extract text from a PDF, preferring PyMuPDF then falling back to pypdf."""
    # Primary: PyMuPDF
    try:
        import pymupdf  # type: ignore

        parts: list[str] = []
        with pymupdf.open(path) as doc:
            for page in doc:
                parts.append(page.get_text("text"))
        text = "\n".join(parts).strip()
        if text:
            return text
    except ImportError:
        try:
            import fitz  # type: ignore  # older PyMuPDF import name

            parts = []
            with fitz.open(path) as doc:
                for page in doc:
                    parts.append(page.get_text("text"))
            text = "\n".join(parts).strip()
            if text:
                return text
        except Exception as exc:  # pragma: no cover - engine specific
            logger.warning("PyMuPDF failed for %s: %s", path, exc)
    except Exception as exc:
        logger.warning("PyMuPDF failed for %s: %s", path, exc)

    # Fallback: pypdf
    try:
        from pypdf import PdfReader

        with open(path, "rb") as handle:
            reader = PdfReader(handle)
            return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception as exc:
        logger.error("pypdf fallback failed for %s: %s", path, exc)
        raise EmptyResumeError("Could not read this PDF. It may be corrupted.") from exc


def _extract_docx(path: str) -> str:
    """Extract text (paragraphs + tables) from a .docx file."""
    import docx

    document = docx.Document(path)
    parts = [para.text for para in document.paragraphs]

    # Tables often hold skills/experience grids - include them.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(part for part in parts if part).strip()


def _extract_plain(path: str) -> str:
    """Read a plain-text/markdown/rtf file with encoding detection."""
    import chardet

    with open(path, "rb") as handle:
        raw = handle.read()

    detected = chardet.detect(raw).get("encoding") or "utf-8"
    try:
        return raw.decode(detected, errors="ignore").strip()
    except (LookupError, UnicodeDecodeError):
        return raw.decode("utf-8", errors="ignore").strip()


def _extract_image(path: str) -> str:
    """OCR an image resume using Tesseract (optional dependency)."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise UnsupportedFileType(
            "Image resumes require pytesseract and the tesseract binary. "
            "Run: sudo pacman -S tesseract tesseract-data-eng"
        ) from exc

    try:
        with Image.open(path) as image:
            return pytesseract.image_to_string(image).strip()
    except Exception as exc:
        logger.error("OCR failed for %s: %s", path, exc)
        raise EmptyResumeError("Could not OCR this image resume.") from exc


def extract_text_from_path(path: str) -> str:
    """Dispatch to the right extractor based on the file extension."""
    extension = os.path.splitext(path)[1].lower()

    if extension == ".pdf":
        text = _extract_pdf(path)
    elif extension == ".docx":
        text = _extract_docx(path)
    elif extension in {".txt", ".md", ".rtf"}:
        text = _extract_plain(path)
    elif extension in {".png", ".jpg", ".jpeg", ".tiff", ".bmp", ".webp"}:
        text = _extract_image(path)
    else:
        raise UnsupportedFileType(
            f"Unsupported file type '{extension}'. "
            f"Allowed: {', '.join(settings.ALLOWED_RESUME_EXTENSIONS)}"
        )

    text = normalise_text(text)
    if not text:
        raise EmptyResumeError(
            "No readable text found in this file. If it is a scanned PDF, "
            "upload an image or a text-based PDF instead."
        )
    return text


def extract_text(uploaded_file) -> str:
    """Extract text from an uploaded file without keeping it on disk."""
    extension = os.path.splitext(uploaded_file.name or "")[1].lower()

    # Text-like files can be read straight from memory.
    if extension in {".txt", ".md"}:
        raw = uploaded_file.read()
        uploaded_file.seek(0)
        import chardet

        detected = chardet.detect(raw).get("encoding") or "utf-8"
        text = raw.decode(detected, errors="ignore")
        text = normalise_text(text)
        if not text:
            raise EmptyResumeError("This file appears to be empty.")
        return text

    suffix = extension or ".bin"
    handle = tempfile.NamedTemporaryFile(suffix=suffix, delete=False)
    temp_path = handle.name
    try:
        for chunk in uploaded_file.chunks():
            handle.write(chunk)
        handle.close()
        return extract_text_from_path(temp_path)
    finally:
        try:
            handle.close()
        except Exception:
            pass
        if os.path.exists(temp_path):
            os.remove(temp_path)


def normalise_text(text: str) -> str:
    """Collapse whitespace and strip control characters for downstream use."""
    if not text:
        return ""
    # Remove NULs and other control chars except newline/tab.
    cleaned = "".join(
        ch for ch in text if ch in "\n\t" or (ch.isprintable() and ch != "\x00")
    )
    lines = [line.strip() for line in cleaned.splitlines()]
    lines = [line for line in lines if line]
    return "\n".join(lines).strip()
