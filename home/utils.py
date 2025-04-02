from django.core.mail import send_mail
from django.conf import settings

# def send_approval_email(company):
#     """
#     Sends an approval email to the company after it gets approved by the admin.
#     """
#     subject = f"Your company {company.name} has been approved!"
#     message = f"Dear {company.name},\n\nWe are happy to inform you that your company has been approved for the platform.\n\nBest regards,\nThe ResumeChecker Team"
    
#     # You can customize the sender's email address
#     from_email = settings.DEFAULT_FROM_EMAIL
#     recipient_list = [company.email]  # Assuming the company has an 'email' field

#     # Send email
#     send_mail(
#         subject,
#         message,
#         from_email,
#         recipient_list,
#         fail_silently=False,
#     )

#     print(f"Approval email sent to {company.name} at {company.email}")



from django.core.mail import send_mail

def send_approval_email(company):
    subject = "Your Company Registration is Approved!"
    message = f"Dear {company.name},\n\nYour company has been approved by the admin. You can now log in and use our platform.\n\nBest Regards,\nResumeChecker Team"
    send_mail(subject, message, "rosemariyaviji@gmail.com", [company.user.email])



import fitz  # PyMuPDF for PDFs
import docx
import os

def parse_pdf(file_path):
    """Extract text from a PDF file using PyMuPDF (fitz)."""
    text = ""
    try:
        with fitz.open(file_path) as doc:
            for page in doc:
                text += page.get_text("text") + "\n"
    except Exception as e:
        print(f"Error reading PDF: {e}")
    return text.strip()

def parse_docx(file_path):
    """Extract text from a DOCX file using python-docx."""
    text = ""
    try:
        doc = docx.Document(file_path)
        text = "\n".join([para.text for para in doc.paragraphs])
    except Exception as e:
        print(f"Error reading DOCX: {e}")
    return text.strip()

def parse_resume(file):
    """Extract text from a resume file (PDF or DOCX)."""
    # Save uploaded file temporarily
    temp_path = f"temp/{file.name}"
    os.makedirs("temp", exist_ok=True)

    with open(temp_path, "wb") as f:
        for chunk in file.chunks():
            f.write(chunk)

    # Determine file type and extract text
    if file.name.lower().endswith(".pdf"):
        extracted_text = parse_pdf(temp_path)
    elif file.name.lower().endswith(".docx"):
        extracted_text = parse_docx(temp_path)
    else:
        extracted_text = ""

    # Clean up temp file
    os.remove(temp_path)
    
    return extracted_text



# import os
# import tempfile
# import PyPDF2
# import docx

# import fitz  # PyMuPDF

# def extract_text_from_pdf(file_path):
#     """Extracts text from a given PDF file."""
#     try:
#         text = ""
#         with fitz.open(file_path) as doc:
#             for page in doc:
#                 text += page.get_text("text") + "\n"

#         return text.strip() if text else "No text found in PDF."
    
#     except Exception as e:
#         print(f"PDF Extraction Error: {e}")
#         return "Error extracting text"

# def extract_text_from_docx(docx_path):
#     doc = docx.Document(docx_path)
#     return "\n".join([para.text for para in doc.paragraphs])

# def parse_resume(file):
#     temp_dir = tempfile.gettempdir()
#     file_name = file.name if hasattr(file, "name") else "temp_resume.pdf"
#     temp_path = os.path.join(temp_dir, file_name)
    
#     # Save the uploaded file to a temporary location
#     with open(temp_path, "wb") as f:
#         for chunk in file.chunks():
#             f.write(chunk)
    
#     # Determine file type and extract text
#     if file_name.endswith(".pdf"):
#         parsed_text = extract_text_from_pdf(temp_path)
#     elif file_name.endswith(".docx"):
#         parsed_text = extract_text_from_docx(temp_path)
#     else:
#         parsed_text = "Unsupported file format. Please upload a PDF or DOCX file."
    
#     return {"parsed_text": parsed_text}




# def recommend_courses(skills):
#     # Example logic to recommend courses based on skills
#     if not skills:
#         return ["No recommendations available"]
    
#     course_map = {
#         "Python": ["Python for Beginners", "Advanced Python"],
#         "Django": ["Django Web Development", "REST API with Django"],
#         "Machine Learning": ["Intro to ML", "Deep Learning with TensorFlow"],
#     }
    
#     recommended = []
#     for skill in skills:
#         recommended.extend(course_map.get(skill, []))
    
#     return recommended or ["No relevant courses found"]




# import fitz  # PyMuPDF

# # Dummy skills dataset
# SKILLS_DATABASE = {
#     "python": ["Python for Beginners", "Advanced Python", "Data Science with Python"],
#     "machine learning": ["ML Basics", "Deep Learning", "AI for Everyone"],
#     "django": ["Django Web Development", "Advanced Django", "Django REST Framework"],
# }

# def extract_text_from_pdf(pdf_path):
#     """Extract text from a PDF using PyMuPDF."""
#     text = ""
#     try:
#         with fitz.open(pdf_path) as doc:
#             for page in doc:
#                 text += page.get_text("text")
#     except Exception as e:
#         return f"Error: {e}"
#     return text.strip() if text else "No text found"

# def extract_skills(text):
#     """Extract skills from text by matching known skills."""
#     found_skills = [skill for skill in SKILLS_DATABASE.keys() if skill.lower() in text.lower()]
#     return found_skills

# def recommend_courses(skills):
#     """Recommend courses based on extracted skills."""
#     courses = []
#     for skill in skills:
#         courses.extend(SKILLS_DATABASE.get(skill.lower(), []))
#     return courses




